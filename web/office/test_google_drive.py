from io import BytesIO, StringIO
from types import SimpleNamespace
from hashlib import sha256
import tempfile
from cryptography.fernet import Fernet
from unittest.mock import MagicMock, patch

from django.contrib.auth.models import User, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import connection
from django.test import TestCase, override_settings
from docx import Document

from .forms import DriveUploadForm, record_form
from .google_drive import create_client_folder, upload_document
from .models import Cliente, Configuration, DocumentoCliente, DriveConnection, DocumentTemplate
from .registry import RESOURCES


class DriveTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser('admin', password='Test-local-2026!')
        self.viewer = User.objects.create_user('viewer')
        self.viewer.user_permissions.add(*Permission.objects.filter(codename__startswith='view_', content_type__app_label='office'))
        self.person = Cliente.objects.create(nome_completo='Ana Teste')
        self.other = Cliente.objects.create(nome_completo='Outro Cliente')
        Configuration.objects.create(category='documento', label='Passaporte')
        self.document = DocumentoCliente.objects.create(cliente=self.person, documento_nome='Passaporte')
        self.client.force_login(self.admin)

    def pdf(self):
        return SimpleUploadedFile('passaporte.pdf', b'%PDF-1.7\nexample', content_type='application/pdf')

    def test_legacy_preflight_does_not_require_new_drive_columns(self):
        from .management.commands.check_legacy_schema import LEGACY
        tables = {model._meta.db_table: model for model in LEGACY}
        def columns(cursor, table):
            return [SimpleNamespace(name=f.column) for f in tables[table]._meta.fields if not f.name.startswith('drive_')]
        with patch.object(connection.introspection, 'table_names', return_value=list(tables)), \
                patch.object(connection.introspection, 'get_table_description', side_effect=columns), \
                patch.object(connection, 'cursor') as cursor:
            cursor.return_value.__enter__.return_value.fetchone.return_value = (0,)
            output = StringIO()
            call_command('check_legacy_schema', stdout=output)
            self.assertIn('Pré-verificação concluída', output.getvalue())

    def test_client_flow_keeps_context_and_ignores_forged_client(self):
        url = f'/gestao/documentos/novo/?cliente={self.person.pk}'
        listing = self.client.get(f'/gestao/documentos/?cliente={self.person.pk}')
        self.assertContains(listing, url)
        form = self.client.get(url)
        self.assertContains(form, 'Ana Teste')
        self.assertNotContains(form, '<select name="cliente"')
        result = self.client.post(url, {'cliente': self.other.pk, 'documento_nome': 'Passaporte'})
        self.assertEqual(result.status_code, 302)
        created = DocumentoCliente.objects.latest('pk')
        self.assertEqual(created.cliente_id, self.person.pk)
        self.assertContains(self.client.get(result.url), f'?cliente={self.person.pk}')

    def test_context_is_retained_on_validation_error_and_cancel(self):
        url = f'/gestao/documentos/novo/?cliente={self.person.pk}'
        result = self.client.post(url, {'documento_nome': ''})
        self.assertEqual(result.status_code, 200)
        self.assertContains(result, f'/gestao/documentos/?cliente={self.person.pk}')
        self.assertEqual(result.context['form'].fields['cliente'].disabled, True)
        self.assertEqual(DocumentoCliente.objects.count(), 1)
        self.assertEqual(self.client.get('/gestao/documentos/novo/?cliente=not-a-client').status_code, 404)
        self.assertEqual(self.client.get('/gestao/documentos/?cliente=99999').status_code, 404)

    def test_documents_require_a_client_context(self):
        self.assertRedirects(self.client.get('/gestao/documentos/'), '/gestao/clientes/')
        self.assertRedirects(self.client.get('/gestao/documentos/novo/'), '/gestao/clientes/')

    def test_upload_permissions_methods_and_oauth_state(self):
        self.client.force_login(self.viewer)
        for url in [f'/documentos/{self.document.pk}/enviar-drive/', f'/clientes/{self.person.pk}/pasta-drive/', '/google/drive/ligar/']:
            self.assertEqual(self.client.post(url, {'ficheiro': self.pdf()}).status_code, 403)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get('/google/drive/ligar/').status_code, 405)
        self.assertEqual(self.client.get('/google/drive/callback/?state=forged&code=forged').status_code, 403)
        self.assertEqual(self.client.get(f'/documentos/{self.document.pk}/enviar-drive/').status_code, 405)

    def test_upload_rejects_unsafe_or_oversize_files(self):
        files = [SimpleUploadedFile('x.html', b'<html>'), SimpleUploadedFile('x.pdf', b'<html>'),
                 SimpleUploadedFile('x.docx', b'not a zip'), SimpleUploadedFile('x.pdf', b'%PDF-' + b'x' * (10 * 1024 * 1024))]
        for file in files:
            with self.subTest(file=file.name):
                self.assertFalse(DriveUploadForm(files={'ficheiro': file}).is_valid())
        buffer = BytesIO()
        Document().save(buffer)
        self.assertTrue(DriveUploadForm(files={'ficheiro': SimpleUploadedFile('x.docx', buffer.getvalue())}).is_valid())
        self.assertTrue(DriveUploadForm(files={'ficheiro': self.pdf()}).is_valid())

    @patch('office.google_drive.upload_document')
    def test_new_record_upload_and_failure_preserve_record(self, upload):
        upload.side_effect = RuntimeError('private token must not leak')
        result = self.client.post(f'/gestao/documentos/novo/?cliente={self.person.pk}', {'documento_nome': 'Passaporte', 'ficheiro': self.pdf()}, follow=True)
        self.assertContains(result, 'o envio ao Drive falhou')
        self.assertNotContains(result, 'private token')
        self.assertEqual(DocumentoCliente.objects.count(), 2)
        self.assertEqual(upload.call_args.args[1], b'%PDF-1.7\nexample')
        self.assertEqual(upload.call_args.args[2], 'passaporte.pdf')

    def files(self):
        files = MagicMock()
        files.list.return_value.execute.return_value = {'files': []}
        files.create.return_value.execute.side_effect = [{'id': 'root-id'}, {'id': 'client-id'}, {'id': 'file-id'}]
        files.get.return_value.execute.side_effect = [{'id': 'root-id'}, {'id': 'client-id', 'name': f'{self.person.pk} - Ana Teste'}]
        return files

    @patch('office.google_drive.get_service')
    def test_real_upload_service_creates_named_folder_and_stores_ids(self, service):
        DriveConnection.objects.create(encrypted_credentials='encrypted-test-only')
        files = self.files()
        service.return_value.files.return_value = files
        result = upload_document(self.document.pk, b'%PDF-test', 'test.pdf', 'application/pdf')
        self.assertEqual(files.create.call_args_list[1].kwargs['body']['name'], f'{self.person.pk} - Ana Teste')
        self.assertEqual(files.create.call_args_list[2].kwargs['body']['parents'], ['client-id'])
        self.assertEqual(result.drive_file_id, 'file-id')
        self.assertTrue(result.entregue)
        self.person.refresh_from_db()
        self.assertEqual(self.person.drive_folder_id, 'client-id')
        with self.assertRaises(ValueError):
            upload_document(self.document.pk, b'%PDF-test', 'again.pdf', 'application/pdf')
        self.assertEqual(files.create.call_count, 3)

    @patch('office.google_drive.get_service')
    def test_retry_recovers_existing_remote_files_without_duplication(self, service):
        DriveConnection.objects.create(encrypted_credentials='encrypted-test-only')
        files = self.files()
        files.list.return_value.execute.side_effect = [{'files': [{'id': 'root-id'}]}, {'files': [{'id': 'client-id'}]}, {'files': [{'id': 'file-id'}]}]
        files.get.return_value.execute.side_effect = [{'id': 'root-id'}, {'id': 'client-id', 'name': f'{self.person.pk} - Ana Teste'},
            {'name': 'test.pdf', 'appProperties': {'himpHash': sha256(b'%PDF-test').hexdigest()}}]
        service.return_value.files.return_value = files
        result = upload_document(self.document.pk, b'%PDF-test', 'test.pdf', 'application/pdf')
        self.assertEqual(result.drive_file_id, 'file-id')
        files.create.assert_not_called()

    @patch('office.google_drive.get_service')
    def test_rename_uses_same_folder_and_remote_failure_rolls_back(self, service):
        DriveConnection.objects.create(encrypted_credentials='encrypted-test-only', root_folder_id='root-id')
        self.person.drive_folder_id = 'client-id'
        self.person.save()
        files = self.files()
        files.get.return_value.execute.side_effect = [{'id': 'root-id'}, {'id': 'client-id', 'name': 'old name'}]
        service.return_value.files.return_value = files
        self.assertEqual(create_client_folder(self.person.pk), 'client-id')
        files.update.assert_called_once_with(fileId='client-id', body={'name': f'{self.person.pk} - Ana Teste'}, fields='id')
        files.get.return_value.execute.side_effect = RuntimeError('offline')
        with self.assertRaises(RuntimeError):
            upload_document(self.document.pk, b'test', 'test.pdf', 'application/pdf')
        self.document.refresh_from_db()
        self.assertFalse(self.document.drive_file_id)
        self.assertFalse(self.document.entregue)

    def test_linked_document_cannot_be_moved_to_another_client(self):
        self.document.drive_file_id = 'file-id'
        self.document.save()
        self.client.post(f'/gestao/documentos/{self.document.pk}/editar/', {'cliente': self.other.pk, 'documento_nome': 'Passaporte'})
        self.document.refresh_from_db()
        self.assertEqual(self.document.cliente_id, self.person.pk)
        self.assertNotIn('drive_folder_id', record_form(RESOURCES['clientes'])().fields)
        self.assertNotIn('drive_file_id', record_form(RESOURCES['documentos'])().fields)

    @patch('office.google_drive.flow')
    def test_callback_encrypts_credentials_and_rejects_partial_consent(self, flow):
        key = Fernet.generate_key()
        credentials = flow.return_value.credentials
        credentials.refresh_token = 'test-refresh'
        credentials.has_scopes.return_value = True
        credentials.to_json.return_value = '{"token":"test-private-token"}'
        def session_state():
            from django.utils import timezone
            session = self.client.session
            session['drive_state'] = 'verified-state'
            session['drive_started_at'] = timezone.now().timestamp()
            session.save()
        with override_settings(TOKEN_ENCRYPTION_KEY=key.decode()):
            session_state()
            response = self.client.get('/google/drive/callback/?state=verified-state&code=test-code')
            self.assertEqual(response.status_code, 302)
            connection = DriveConnection.objects.get()
            self.assertNotIn('test-private-token', connection.encrypted_credentials)
            self.assertEqual(Fernet(key).decrypt(connection.encrypted_credentials.encode()), b'{"token":"test-private-token"}')
            credentials.has_scopes.return_value = False
            session_state()
            self.client.get('/google/drive/callback/?state=verified-state&code=test-code')
            connection.refresh_from_db()
            self.assertEqual(Fernet(key).decrypt(connection.encrypted_credentials.encode()), b'{"token":"test-private-token"}')

    @patch('office.google_drive.upload_document')
    def test_generated_document_can_be_saved_and_failure_offers_download(self, upload):
        with tempfile.TemporaryDirectory() as temp, override_settings(MEDIA_ROOT=temp,
                STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
                          'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}}):
            buffer = BytesIO()
            doc = Document()
            doc.add_paragraph('Olá {{ nome }}')
            doc.save(buffer)
            template = DocumentTemplate.objects.create(name='Modelo', file=SimpleUploadedFile('modelo.docx', buffer.getvalue()))
            url = f'/gerar-documento/?cliente={self.person.pk}'
            response = self.client.post(url, {'cliente': self.other.pk, 'modelo': template.pk, 'guardar_drive': 'on'})
            self.assertEqual(response.status_code, 302)
            saved = DocumentoCliente.objects.latest('pk')
            self.assertEqual(saved.cliente_id, self.person.pk)
            self.assertEqual(upload.call_args.args[0], saved.pk)
            self.assertTrue(upload.call_args.args[1].startswith(b'PK'))
            upload.side_effect = RuntimeError('offline')
            response = self.client.post(url, {'modelo': template.pk, 'guardar_drive': 'on'})
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, 'O envio ao Drive falhou')
            self.assertContains(response, 'Abrir registo do documento')
            response = self.client.post(url, {'modelo': template.pk})
            self.assertTrue(response['Content-Type'].startswith('application/vnd.openxmlformats'))
            output = BytesIO(b''.join(response.streaming_content))
            self.assertEqual(Document(output).paragraphs[0].text, 'Olá Ana Teste')
            self.assertEqual(DocumentoCliente.objects.count(), 3)
