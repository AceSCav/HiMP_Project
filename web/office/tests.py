import tempfile
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch, MagicMock

from django.contrib.auth.models import User, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document

from .forms import record_form, TemplateAdminForm
from .models import Cliente, Processo, Pagamento, Agendamento, AuditEvent, Configuration, DocumentTemplate
from .registry import RESOURCES


@override_settings(STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'}, 'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class OfficeTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser('responsavel', 'admin@example.invalid', 'Senha-segura-para-testes-2026')
        self.viewer = User.objects.create_user('consulta', password='Senha-segura-para-testes-2026')
        self.viewer.user_permissions.add(*Permission.objects.filter(content_type__app_label='office', codename__startswith='view_'))
        self.client_record = Cliente.objects.create(nome_completo='Cliente de Teste', nif='123456789')

    def login(self, user=None):
        self.client.force_login(user or self.admin)

    def test_anonymous_cannot_read_records(self):
        for url in ['/', '/gestao/clientes/', '/gestao/clientes/1/', '/gerar-documento/']:
            self.assertEqual(self.client.get(url).status_code, 302)

    def test_every_resource_renders_and_form_renders(self):
        self.login()
        for resource in RESOURCES:
            with self.subTest(resource=resource):
                self.assertEqual(self.client.get(reverse('record_list', args=[resource])).status_code, 200)
                self.assertEqual(self.client.get(reverse('record_create', args=[resource])).status_code, 200)
        self.assertEqual(self.client.get('/').status_code, 200)

    def test_read_only_user_cannot_mutate_or_manage_configuration(self):
        self.login(self.viewer)
        self.assertEqual(self.client.get('/gestao/clientes/').status_code, 200)
        self.assertEqual(self.client.post('/gestao/clientes/novo/', {'nome_completo': 'Intrusão'}).status_code, 403)
        self.assertEqual(self.client.post(f'/gestao/clientes/{self.client_record.pk}/editar/', {'nome_completo': 'Intrusão'}).status_code, 403)
        self.assertEqual(self.client.post(f'/gestao/clientes/{self.client_record.pk}/eliminar/').status_code, 403)
        self.assertEqual(self.client.post('/google/ligar/').status_code, 403)
        self.assertEqual(self.client.get('/gerar-documento/').status_code, 403)
        self.assertEqual(Cliente.objects.get(pk=self.client_record.pk).nome_completo, 'Cliente de Teste')

    def test_authenticated_without_permissions_sees_no_data(self):
        user = User.objects.create_user('sem-acesso', password='Senha-segura-para-testes-2026')
        self.login(user)
        response = self.client.get('/')
        self.assertNotContains(response, 'Cliente de Teste')
        self.assertEqual(self.client.get('/gestao/clientes/').status_code, 403)

    def test_csrf_enforced_and_delete_get_is_safe(self):
        secure = Client(enforce_csrf_checks=True)
        secure.force_login(self.admin)
        self.assertEqual(secure.post('/gestao/clientes/novo/', {'nome_completo': 'Sem token'}).status_code, 403)
        self.assertEqual(secure.get(f'/gestao/clientes/{self.client_record.pk}/eliminar/').status_code, 200)
        self.assertTrue(Cliente.objects.filter(pk=self.client_record.pk).exists())

    def test_password_hashed_and_invalid_login_generic(self):
        self.assertNotEqual(self.admin.password, 'Senha-segura-para-testes-2026')
        self.assertTrue(self.admin.check_password('Senha-segura-para-testes-2026'))
        response = self.client.post('/entrar/', {'username': 'responsavel', 'password': 'wrong'})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_login_rate_limit(self):
        for _ in range(5):
            self.client.post('/entrar/', {'username': 'responsavel', 'password': 'wrong'})
        response = self.client.post('/entrar/', {'username': 'responsavel', 'password': 'Senha-segura-para-testes-2026'})
        self.assertEqual(response.status_code, 429)
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_dynamic_catalogue_changes_available_without_code_change(self):
        Configuration.objects.create(category='estado_civil', label='Novo estado de teste')
        form = record_form(RESOURCES['clientes'])()
        self.assertIn(('Novo estado de teste', 'Novo estado de teste'), form.fields['estado_civil'].choices)
        config = Configuration.objects.get(label='Novo estado de teste')
        config.active = False
        config.save()
        self.assertNotIn(('Novo estado de teste', 'Novo estado de teste'), record_form(RESOURCES['clientes'])().fields['estado_civil'].choices)
        self.client_record.estado_civil = 'Novo estado de teste'
        self.client_record.save()
        self.assertIn(('Novo estado de teste', 'Novo estado de teste'), record_form(RESOURCES['clientes'])(instance=self.client_record).fields['estado_civil'].choices)

    def test_client_create_update_search_and_audit(self):
        self.login()
        response = self.client.post('/gestao/clientes/novo/', {'nome_completo': 'Pessoa Nova', 'email': 'nova@example.invalid'})
        self.assertEqual(response.status_code, 302)
        record = Cliente.objects.get(nome_completo='Pessoa Nova')
        self.assertIsNone(record.nif)
        response = self.client.post(f'/gestao/clientes/{record.pk}/editar/', {'nome_completo': 'Pessoa Atualizada', 'email': 'nova@example.invalid'})
        self.assertEqual(response.status_code, 302)
        self.assertContains(self.client.get('/gestao/clientes/?q=Atualizada'), 'Pessoa Atualizada')
        self.assertEqual(AuditEvent.objects.filter(object_id=str(record.pk)).count(), 2)

    def test_duplicate_identifiers_and_date_validation(self):
        form = record_form(RESOURCES['clientes'])({'nome_completo': 'Outro cliente', 'nif': '123456789'})
        self.assertFalse(form.is_valid())
        form = record_form(RESOURCES['clientes'])({'nome_completo': 'Outro cliente', 'nif': '123', 'data_nascimento': '2999-01-01'})
        self.assertFalse(form.is_valid())
        self.assertIn('nif', form.errors)
        self.assertIn('data_nascimento', form.errors)

    def test_financial_precision_and_reference_zeroes(self):
        self.login()
        response = self.client.post('/gestao/pagamentos/novo/', {'cliente': self.client_record.pk, 'entidade': '00012', 'referencia': '000123456', 'montante': '19.99', 'data_limite': '2026-01-01'})
        self.assertEqual(response.status_code, 302)
        payment = Pagamento.objects.get()
        self.assertEqual(payment.montante, Decimal('19.99'))
        self.assertEqual(payment.referencia, '000123456')
        self.assertEqual(payment.entidade, '00012')
        self.assertContains(self.client.get('/gestao/pagamentos/'), '19,99')

    def test_negative_amount_rejected(self):
        form = record_form(RESOURCES['pagamentos'])({'cliente': self.client_record.pk, 'montante': '-1'})
        self.assertFalse(form.is_valid())

    def test_linked_client_deletion_blocked_and_audit_rolled_back(self):
        self.login()
        Processo.objects.create(cliente=self.client_record, numero_processo='TESTE-1')
        self.client.post(f'/gestao/clientes/{self.client_record.pk}/eliminar/')
        self.assertTrue(Cliente.objects.filter(pk=self.client_record.pk).exists())
        self.assertFalse(AuditEvent.objects.filter(action='delete').exists())

    def test_xss_escaped_and_private_cache_headers(self):
        self.client_record.nome_completo = '<script>alert(1)</script>'
        self.client_record.save()
        self.login()
        response = self.client.get('/gestao/clientes/')
        self.assertNotContains(response, '<script>alert(1)</script>')
        self.assertContains(response, '&lt;script&gt;')
        self.assertEqual(response['Cache-Control'], 'no-store, private')
        self.assertIn("frame-ancestors 'none'", response['Content-Security-Policy'])

    def test_document_generation_escapes_xml_and_requires_permissions(self):
        self.login()
        with tempfile.TemporaryDirectory() as temp, override_settings(MEDIA_ROOT=temp):
            doc = Document()
            doc.add_paragraph('Cliente: {{ nome }}; NIF: {{ nif }}')
            buffer = BytesIO()
            doc.save(buffer)
            template = DocumentTemplate.objects.create(name='Modelo teste', file=SimpleUploadedFile('modelo.docx', buffer.getvalue()), required_fields=['nome', 'nif'])
            self.client_record.nome_completo = 'Pessoa & Filhos <Teste>'
            self.client_record.save()
            response = self.client.post('/gerar-documento/', {'cliente': self.client_record.pk, 'modelo': template.pk})
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.streaming)
            result = Document(BytesIO(b''.join(response.streaming_content)))
            self.assertIn('Pessoa & Filhos <Teste>', result.paragraphs[0].text)
            self.assertIn('attachment;', response['Content-Disposition'])

    def test_invalid_template_and_missing_fields_rejected(self):
        form = TemplateAdminForm(data={'name': 'Inválido', 'required_fields': '["nome"]', 'active': True}, files={'file': SimpleUploadedFile('fake.docx', b'not-a-zip')})
        self.assertFalse(form.is_valid())
        self.assertIn('file', form.errors)

    def test_oauth_state_rejection_and_connect_get_safe(self):
        self.login()
        self.assertEqual(self.client.get('/google/ligar/').status_code, 405)
        self.assertEqual(self.client.get('/google/callback/?state=forged&code=forged').status_code, 403)

    def test_sync_failure_preserves_local_record(self):
        self.login()
        event = Agendamento.objects.create(cliente=self.client_record, titulo='Teste', data_inicio=timezone.now(), duracao='30 minutos')
        with patch('office.google_calendar.sync_event', side_effect=RuntimeError('private error')):
            response = self.client.post(f'/agenda/{event.pk}/sincronizar/', follow=True)
        self.assertContains(response, 'A sincronização falhou')
        self.assertNotContains(response, 'private error')
        self.assertTrue(Agendamento.objects.filter(pk=event.pk).exists())

    def test_all_core_records_create_edit_delete(self):
        self.login()
        data = {
            'processos': {'cliente': self.client_record.pk, 'numero_processo': 'TESTE-2'},
            'agenda': {'cliente': self.client_record.pk, 'titulo': 'Consulta teste', 'data_inicio': '2026-10-15T14:00', 'duracao': '30 minutos', 'motivo': 'Consulta'},
            'documentos': {'cliente': self.client_record.pk, 'documento_nome': Configuration.objects.filter(category='documento').first().label, 'entregue': 'on'},
        }
        for resource, payload in data.items():
            with self.subTest(resource=resource):
                response = self.client.post(reverse('record_create', args=[resource]), payload)
                self.assertEqual(response.status_code, 302)
                record = RESOURCES[resource].model.objects.get()
                self.assertEqual(self.client.get(reverse('record_detail', args=[resource, record.pk])).status_code, 200)
                self.assertEqual(self.client.post(reverse('record_edit', args=[resource, record.pk]), payload).status_code, 302)
                self.assertEqual(self.client.post(reverse('record_delete', args=[resource, record.pk])).status_code, 302)
                self.assertFalse(RESOURCES[resource].model.objects.exists())

    def test_step_creation_and_case_history(self):
        from .models import Etapa, Fase
        self.login()
        process = Processo.objects.create(cliente=self.client_record, numero_processo='ETAPA-1')
        phase = Fase.objects.create(fase='Análise')
        payload = {'processo': process.pk, 'fase': phase.pk, 'data_fase': '2026-10-07', 'observacao': 'Recebido'}
        self.assertEqual(self.client.post('/gestao/etapas/novo/', payload).status_code, 302)
        step = Etapa.objects.get()
        self.assertContains(self.client.get(f'/gestao/processos/{process.pk}/'), 'Análise')
        self.assertEqual(self.client.post(f'/gestao/etapas/{step.pk}/editar/', payload).status_code, 302)
        self.assertEqual(self.client.post(f'/gestao/etapas/{step.pk}/eliminar/').status_code, 302)

    def test_google_sync_uses_minimal_scope_and_stable_event_id(self):
        from .google_calendar import sync_event, SCOPES
        from .models import CalendarConnection
        from cryptography.fernet import Fernet
        import json
        from googleapiclient.errors import HttpError
        import httplib2
        key = Fernet.generate_key()
        token = {'token': 'test-token', 'refresh_token': 'test-refresh', 'token_uri': 'https://oauth2.googleapis.com/token', 'client_id': 'test', 'client_secret': 'test', 'expiry': '2999-01-01T00:00:00Z'}
        CalendarConnection.objects.create(encrypted_credentials=Fernet(key).encrypt(json.dumps(token).encode()).decode())
        event = Agendamento.objects.create(titulo='Teste', data_inicio=timezone.now(), duracao='30 minutos')
        events = MagicMock()
        events.get.return_value.execute.side_effect = HttpError(httplib2.Response({'status': '404'}), b'{}')
        events.insert.return_value.execute.return_value = {'id': f'himp{event.pk:x}'}
        service = MagicMock()
        service.events.return_value = events
        with override_settings(TOKEN_ENCRYPTION_KEY=key.decode()), patch('office.google_calendar.build', return_value=service):
            sync_event(event)
        self.assertEqual(event.google_event_id, f'himp{event.pk:x}')
        self.assertEqual(events.insert.call_args.kwargs['body']['id'], event.google_event_id)
        self.assertEqual(SCOPES, ['https://www.googleapis.com/auth/calendar.events'])
